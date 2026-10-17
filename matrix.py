"""Matriks sintesis literatur (tabel telaah jurnal) dari koleksi user.

Kolom: No | Penulis (Tahun) | Judul | Jurnal | Metode | Hasil kunci.
Metode & hasil diekstrak best-effort dari abstrak (ID/EN).
"""
import logging
import re
from io import BytesIO

log = logging.getLogger(__name__)

_METHOD_KEYS = ("method", "design", "sample", "participants", "respondent",
                "metode", "desain", "sampel", "responden", "partisipan",
                "instrument", "instrumen", "data were", "data di")
_RESULT_KEYS = ("result", "found", "showed", "significant", "conclud",
                "hasil", "menunjukkan", "ditemukan", "signifikan", "simpulan",
                "p <", "p<", "p =", "p=")


def _sentences(text):
    text = re.sub(r"\s+", " ", text or "").strip()
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if len(s.strip()) > 25]


def _pick(sents, keys, max_chars=280):
    hits = [s for s in sents if any(k in s.lower() for k in keys)]
    out, total = [], 0
    for s in hits[:3]:
        if total + len(s) > max_chars:
            break
        out.append(s)
        total += len(s)
    return " ".join(out)


def extract_method_result(abstract):
    """Return (metode, hasil) best-effort dari abstrak."""
    sents = _sentences(abstract)
    metode = _pick(sents, _METHOD_KEYS)
    hasil = _pick([s for s in sents if s not in metode], _RESULT_KEYS)
    return metode, hasil


def _authors_short(authors):
    if not authors:
        return "-"
    first = authors[0]
    fam = first.split(",")[0].strip() if "," in first else first.strip().split()[-1]
    return fam + (" et al." if len(authors) > 1 else "")


def build_matrix_docx(items):
    """Bangun file .docx tabel matriks dari list work. Return bytes."""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    doc.add_heading("Matriks Sintesis Literatur", level=1)
    p = doc.add_paragraph()
    r = p.add_run(f"{len(items)} referensi — diekstrak otomatis dari koleksi. "
                  "Kolom Metode & Hasil diambil best-effort dari abstrak; "
                  "periksa dan lengkapi manual sebelum dipakai di skripsi.")
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    headers = ["No", "Penulis (Tahun)", "Judul", "Jurnal", "Metode", "Hasil Kunci"]
    widths = [Cm(1), Cm(3), Cm(5), Cm(3.5), Cm(5), Cm(5)]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        cell.width = widths[i]

    for n, w in enumerate(items, 1):
        row = table.add_row()
        metode, hasil = extract_method_result(w.get("abstract"))
        vals = [
            str(n),
            f"{_authors_short(w.get('authors'))} ({w.get('year') or '-'})",
            w.get("title") or "-",
            w.get("journal") or "-",
            metode or "-",
            hasil or "-",
        ]
        for i, v in enumerate(vals):
            cell = row.cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            cell.width = widths[i]

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
